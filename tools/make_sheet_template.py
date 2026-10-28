"""Regenerates sheet-template/WhatsApp-Bot-Sheet-Template.xlsx.

Run after changing any Sheet column, FAQ example or default message:
    python tools/make_sheet_template.py
"""
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from bot import DEFAULT_MESSAGES

F = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E78")
EXAMPLE = PatternFill("solid", fgColor="FFF2CC")
wb = Workbook()

def table(ws, headers, rows, widths, notes, example_fill=True):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c, w in zip(ws[1], widths):
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[c.column_letter].width = w
        if c.value in notes:
            c.comment = Comment(notes[c.value], "Setup guide", width=320, height=160)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if example_fill:
                c.fill = EXAMPLE
    ws.freeze_panes = "A2"

# --- Read me ---------------------------------------------------------------
rm = wb.active
rm.title = "Read me"
lines = [
    ("WhatsApp lead bot - control panel", True),
    ("The bot reads Sites, FAQ and Messages (changes show up within a minute) and writes every enquiry to Leads.", False),
    ("", False),
    ("What to edit", True),
    ("Sites - one row per project. Yellow rows are EXAMPLES: replace them with your projects before testing.", False),
    ("FAQ - questions the bot answers on its own. Yellow rows are examples: fill in the [brackets].", False),
    ("Messages - the bot's wording. Edit the 'text' column only; leave 'key' exactly as it is.", False),
    ("Leads - written by the bot. Don't edit or reorder its columns; add your own columns to the right if you like.", False),
    ("", False),
    ("Rules", True),
    ("Keep the tab names and the header row (row 1) of every tab exactly as they are.", False),
    ("Sites > active: Yes shows the project in the menu, No hides it. WhatsApp menus show at most 10 projects.", False),
    ("Sites > tagline: shown under the project name in the menu, up to 72 characters.", False),
    ("Sites > brochure_url: a Google Drive link set to 'Anyone with the link', or a direct link to a PDF on your website. WhatsApp accepts PDFs up to 100 MB; smaller files open faster on mobile data.", False),
    ("Sites > maps_link: open the site in Google Maps, tap Share, Copy link, paste it here. The bot turns it into a map pin.", False),
    ("Sites > latitude / longitude: optional. Fill them only if the pin lands in the wrong place; they override maps_link.", False),
    ("Sites > sales_phone: the WhatsApp number customers are handed to for this project. Blank uses the default in the bot's settings.", False),
    ("Sites > alert_phones: everyone who should get a WhatsApp alert about this project's enquiries, comma-separated (3-4 salespeople is fine). Blank alerts sales_phone only.", False),
    ("FAQ > keywords: comma-separated words or phrases. Add spellings customers actually use (rate, rates, price, cost).", False),
    ("FAQ > answer: write just SALES to have a salesperson handle these questions instead of the bot (e.g. price, rate, carpet area). The customer gets the hand-off message and contact card, and the sales team is alerted. A SALES row wins over every other row that matches.", False),
    ("FAQ > site_id: blank = applies to every project; a site_id = only for that project (it wins over a general answer).", False),
    ("AI answers (if switched on): questions the keywords miss are answered by AI using ONLY the FAQ and Sites tabs, else handed to sales. SALES topics always go to a person. Check Leads rows with event ai_answered now and then.", False),
    ("Placeholders you can use in FAQ answers and Messages: {name} {company} {site_name} {availability} {sales_name} {sales_phone}", False),
    ("Leads alerts: Tools > Notification settings > Edit notifications > 'Any changes are made' > 'Right away' emails you on every new lead.", False),
]
for i, (t, bold) in enumerate(lines, start=1):
    c = rm.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=bold, size=14 if i == 1 else 11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
rm.column_dimensions["A"].width = 120

# --- Sites -----------------------------------------------------------------
sites = wb.create_sheet("Sites")
table(
    sites,
    ["site_id", "name", "aliases", "active", "tagline", "brochure_url", "brochure_filename", "availability",
     "maps_link", "latitude", "longitude", "address", "sales_name", "sales_phone", "alert_phones"],
    [
        ["GV", "Green Valley Residency", "green valley, gv", "Yes", "2 & 3 BHK · Possession Dec 2027",
         "https://drive.google.com/file/d/REPLACE_WITH_YOUR_FILE_ID/view?usp=sharing", "Green Valley Residency brochure.pdf",
         "2 BHK (1,050 sq ft): 5 flats left\n3 BHK (1,420 sq ft): 7 flats left", "https://maps.app.goo.gl/REPLACE_WITH_YOUR_LINK", None, None,
         "Plot 12, Example Road, Your City 431001", "Rahul (Sales)", "+91 98765 43210", "+91 98765 43210, +91 98111 22233, +91 99000 44455"],
        ["SH", "Sunrise Heights", "sunrise", "Yes", "3 BHK premium · Ready to move",
         "https://drive.google.com/file/d/REPLACE_WITH_YOUR_FILE_ID/view?usp=sharing", "Sunrise Heights brochure.pdf",
         "3 BHK (1,650 sq ft): 2 flats left", "https://maps.app.goo.gl/REPLACE_WITH_YOUR_LINK", None, None,
         "Survey No. 45, Example Nagar, Your City 431005", "Sneha (Sales)", "+91 91234 56789", "+91 91234 56789, +91 98111 22233"],
    ],
    [9, 26, 20, 8, 34, 48, 30, 40, 42, 11, 11, 36, 18, 18, 52],
    {
        "site_id": "Short unique code, no spaces (GV, SH, PH2). Don't change it once leads exist.",
        "aliases": "Other names customers type for this project, comma-separated. Optional.",
        "active": "Yes or No.",
        "tagline": "Shown under the name in the menu. Max 72 characters.",
        "brochure_url": "Google Drive share link (Anyone with the link) or direct PDF link.",
        "availability": "Shown above the buttons after a customer picks this project. Update it whenever a flat sells.",
        "maps_link": "Google Maps > the site > Share > Copy link. The bot reads the pin from it.",
        "latitude": "Optional. Fill only if the pin from maps_link lands in the wrong place.",
        "longitude": "Optional. Fill only if the pin from maps_link lands in the wrong place.",
        "sales_phone": "With country code, e.g. +91 98765 43210. Blank = default salesperson.",
        "alert_phones": "Everyone who gets a WhatsApp alert for this project's enquiries, comma-separated. Blank = sales_phone only. Needs the alert template (see the guide).",
    },
)
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
sites.add_data_validation(dv)
dv.add("D2:D200")

# --- FAQ -------------------------------------------------------------------
faq = wb.create_sheet("FAQ")
table(
    faq,
    ["keywords", "answer", "site_id"],
    [
        ["price, prices, cost, rate, rates, budget",
         "{site_name}: 2 BHK from Rs [__] lakh, 3 BHK from Rs [__] lakh (all-inclusive except registration). Exact quotes depend on floor and facing.", "GV"],
        ["price, prices, cost, rate, rates, budget",
         "{site_name}: 3 BHK from Rs [__] lakh. Exact quotes depend on floor and facing.", "SH"],
        ["possession, handover, ready, completion",
         "{site_name}: possession is planned for [month year].", "GV"],
        ["rera, rera number, registration number",
         "{site_name} is RERA registered: [RERA number].", "GV"],
        ["rera, approved, approval",
         "All our projects are RERA registered. Pick a project from the menu for its RERA number.", ""],
        ["loan, home loan, bank, emi, finance",
         "Home loans are available from [bank names]. Our team helps with the paperwork.", ""],
        ["payment plan, payment, booking amount, down payment",
         "Booking amount is Rs [__]. The rest is paid in construction-linked instalments.", ""],
        ["amenities, gym, pool, swimming, clubhouse, garden, lift",
         "{site_name} has [amenities].", "GV"],
        ["carpet area, built up, sq ft, sqft, square feet, floor plan",
         "SALES", ""],
        ["parking, car parking",
         "Every flat comes with [one covered car parking].", ""],
        ["maintenance, society charges",
         "Maintenance is about Rs [__] per sq ft per month after possession.", ""],
        ["site visit, visit, timing, timings, office hours, sunday",
         "Our site office is open [10 am to 7 pm, all days]. Tap Book site visit after choosing a project and we'll call you to fix a time.", ""],
    ],
    [42, 90, 10],
    {
        "keywords": "Comma-separated words/phrases. Whole words only: 'rate' does not match 'rates', so list both.",
        "answer": "Reply text. Can use {site_name} {availability} {sales_name} {sales_phone}. "
                  "Or write just SALES: the customer is handed to a salesperson instead.",
        "site_id": "Blank = all projects. A site_id = only that project, and it wins over a general answer.",
    },
)

# --- Messages ----------------------------------------------------------------
used_when = {
    "welcome": "First message with the project menu",
    "menu_button": "Label of the menu button (max 20 characters)",
    "menu_footer": "Small grey line under the menu (max 60 characters)",
    "site_details": "Text under the brochure after a project is picked",
    "btn_visit": "Button 1 after a project is picked (max 20 characters)",
    "btn_price": "Button 2 (max 20 characters); looks for a FAQ row with the keyword 'price'",
    "btn_sales": "Button 3 (max 20 characters)",
    "visit_reply": "Reply to Book site visit; the salesperson's contact card follows",
    "handoff": "Reply when the bot can't answer; the salesperson's contact card follows",
    "faq_footer": "Line added under every FAQ answer",
    "non_text": "Reply to voice notes, photos and other non-text messages",
    "no_sites": "Reply when no project is marked active",
}
msgs = wb.create_sheet("Messages")
table(msgs, ["key", "text", "used_when"],
      [[k, v, used_when[k]] for k, v in DEFAULT_MESSAGES.items()],
      [16, 80, 52],
      {"key": "Don't change these.", "text": "Edit freely. *bold* and line breaks work in WhatsApp."},
      example_fill=False)
for row in msgs.iter_rows(min_row=2, max_col=1):
    row[0].font = Font(name=F, color="808080")

# --- Leads -------------------------------------------------------------------
leads = wb.create_sheet("Leads")
table(leads, ["timestamp", "phone", "name", "site_id", "site_name", "event", "detail"], [],
      [20, 16, 20, 9, 26, 16, 60],
      {"event": "menu_shown, site_selected, faq_answered, ai_answered, price_request, visit_request, sales_request, handoff, non_text",
       "detail": "What the customer typed or tapped."})

out = ROOT / "sheet-template" / "WhatsApp-Bot-Sheet-Template.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print("saved", out)
